from src.utils import paths_utils as paths
from openpyxl.reader.excel import load_workbook
from src.utils.dict_result import DictResult
from src.mapper.local.local_db_insert_impl import LocalMysqlDatabase as LocalInsertMysqlDatabase

def public_batch_insert_mediate_upgrade(data):
    # 获取文件路径
    data_dir = paths.get_resource_path("resources/insert_mediate_upgrade")

    # 读取表格
    workbook = load_workbook(f"{data_dir}/{data['mediateUpgradeFileName']}.xlsx")
    sheet = workbook['Sheet1']

    rows_data_list = list(sheet.iter_rows(min_row=4, max_row=sheet.max_row))
    mediate_upgrade_list = [
        {
            # 原客户名称
            'old_company_name': row[1].value,
            # 原到账时间
            'old_accounting_date': row[2].value,
            # 愿合同开始时间
            'old_start_time': row[3].value,
            # 原合同结束时间
            'old_end_time': row[4].value,
            # 原收款金额
            'old_receive_money': row[5].value,
            # 原质检审核时间
            'old_quality_audit': row[6].value,
            # 原备注
            'old_remark': row[7].value,
            # 新客户名称
            'new_company_name': row[8].value,
            # 新到账时间
            'new_accounting_date': row[9].value,
            # 新到账金额
            'new_receive_money': row[10].value,
            # 新套餐类型
            'new_criterion_type': row[11].value,
            # 新质检审核时间
            'new_quality_audit': row[12].value,
            # 新备注
            'new_remark': row[13].value
        } for row in rows_data_list if row[1].value is not None
    ]

    # 批量新增
    LocalInsertMysqlDatabase().batch_insert_mediate_upgrade(mediate_upgrade_list)

    return DictResult.success_response(
        message=f"新增{len(mediate_upgrade_list)}条调解升级客户成功",
        status_code=200,
        types="batch_insert_mediate_upgrade",
        data=None
    )
