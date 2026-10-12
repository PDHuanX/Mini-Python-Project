import datetime
import src.utils.paths_utils as paths_utils
from openpyxl.reader.excel import load_workbook
from openpyxl import Workbook
from src.utils.dict_result import DictResult
from src.utils.sheet_column_index_utils import ColumnIndexTow

'''对外接口'''
def public_get_server_during_data(data):
    return _get_server_during_data(data)

'''获取服务顾问在期数'''
def _get_server_during_data(data):
    # 获取截止时间，因为有可能最后一个工作日不是月度最后一天
    deadline = datetime.datetime.strptime(data['deadline'], '%Y-%m-%d')

    # 拆解需要的服务顾问
    server_string_list = data['serverStr'].split(',') if ',' in data['serverStr'] else data['serverStr'].split('，')

    # 获取文件路径
    data_sir = paths_utils.get_resource_path('resources/excel_templates')

    # 获取表格
    workbook = load_workbook(f'{data_sir}/{data['sourceDataFileNameServer']}.xlsx')
    sheet = workbook['线下客户表']

    # 获取表格数据
    sheet_row_data_list = list(sheet.iter_rows(min_row=2, max_row=sheet.max_row))

    # 结果集
    result_list = {}

    for server in server_string_list:
        for row in sheet_row_data_list:
            # 服务顾问
            server_name = row[ColumnIndexTow.SERVE_NAME].value
            if server_name is None:
                continue

            # 获取 客户名称单元格的颜色，过滤掉不需要的
            fg_color = row[ColumnIndexTow.COMPANY_NAME].fill.fgColor
            # 如果是这些颜色的数据，跳过，这些数据不算户数
            if fg_color.rgb in ['FFFFFF00', 'FFC0C0C0', 'FF666699', '00C0C0C0']:
                continue

            # 合同结束时间
            end_time = row[ColumnIndexTow.END_TIME].value
            # 如果到期时间为空，说明不是会员或调解套餐
            if end_time is None or not end_time: continue
            # 只取还在期的客户
            if end_time <= deadline: continue

            if server == server_name:
                # 如果不存在这个key，初始化
                if server not in result_list.keys():
                    result_list[server] = {'customers': [], 'count': 0}

                # id集，用于数据出错时溯源
                result_list[server]['customers'].append((
                    row[ColumnIndexTow.CUSTOMER_ID].value,
                    row[ColumnIndexTow.COMPANY_NAME].value,
                    row[ColumnIndexTow.START_TIME].value,
                    row[ColumnIndexTow.END_TIME].value,
                    row[ColumnIndexTow.MEMBER_TYPE].value,
                    row[ColumnIndexTow.SERVE_NAME].value,
                    row[ColumnIndexTow.LAWERS].value
                ))
                result_list[server]['count'] += 1

    # 创建新表
    new_workbook = Workbook()
    new_sheet_one = new_workbook['Sheet']
    new_sheet_one.title = '汇总'
    # 设置表头
    new_sheet_one.append(['顾问名称', '在期户数'])
    data_list = [
        (key, val['count']) for key, val in result_list.items()
    ]
    # 设置数据
    for row in data_list:
        new_sheet_one.append(row)

    # 表头
    handers = ['客户id', '客户名称', '合同开始时间', '合同结束时间', '套餐类型', '服务顾问', '法律顾问']

    # 循环处理数据
    for key, val in result_list.items():
        # 根据顾问名称生成新子表
        new_sheet = new_workbook.create_sheet(key)
        # 设置表头
        new_sheet.append(handers)

        # 这个val本身也是dict类型，直接取customers即可
        for customer in val['customers']:
            new_sheet.append(customer)

    # 保存表格
    new_workbook.save(f'{paths_utils.get_resource_path('resources/server_during')}/{deadline.strftime("%Y-%m-%d")}服务顾问在期户数.xlsx')

    return DictResult.success_response(
        message="获取成功，请前往【resources/server_during】文件夹查看",
        status_code=200,
        types="submit_get_server_during_data",
        data=None
    )