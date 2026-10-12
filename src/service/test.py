from src.repository.online.db_query_impl import MysqlDatabase as QueryMysqlDatabase
from src.utils import paths_utils as paths
from openpyxl import load_workbook
from openpyxl import Workbook
from tqdm import tqdm

def test1():
    data_dir = paths.get_resource_path("resources/test")
    workbook = load_workbook(f"{data_dir}/5月计提.xlsx")
    sheet = workbook["Sheet1"]

    final_result = []

    service_result = QueryMysqlDatabase().query_test3()

    row_list = list(sheet.iter_rows())
    for row in tqdm(row_list, desc="获取信息中......"):
        law_name = row[0].value
        salesman_name = row[1].value
        company_name = row[2].value
        service_item = row[3].value
        is_end = row[4].value
        is_accrue = row[5].value
        service_name = row[6].value

        json = {
            "customer_id": "",
            "law_name": law_name,
            "salesman_name": salesman_name,
            "company_name": company_name,
            "service_item": service_item,
            "is_end": is_end,
            "is_accrue": is_accrue,
            "service_name": service_name,
            "start_time": "",
            "end_time": ""
        }

        for service in service_result:

            if (service['lawers'] == law_name or service['consultant_name'] == law_name) and (service['company_name'] in company_name or service['company_name'] in company_name + "（调解升级）"):
                json['customer_id'] = service["customer_id"]
                json['start_time'] = service["start_time"]
                json['end_time'] = service["end_time"]

        final_result.append(json)



    # 创建新表
    new_workbook = Workbook()
    # 汇总表格
    new_sheet_one = new_workbook['Sheet']
    new_sheet_one.title = '汇总'
    new_sheet_one.append(
        ['客户ID', '姓名', '服务城市', '客户名称', '今日完成事项', '是否完结', '是否计提', '合同类型备注', '开始时间', '结束时间'])

    for row in final_result:
        row_data = [
            row["customer_id"],
            row["law_name"],
            row["salesman_name"],
            row["company_name"],
            row["service_item"],
            row["is_end"],
            row["is_accrue"],
            row["service_name"],
            row["start_time"],
            row["end_time"]
        ]
        new_sheet_one.append(row_data)

    new_workbook.save(
        f'{paths.get_resource_path('resources/test')}/立案事项数据4.xlsx')

def test():
    data_dir = paths.get_resource_path("resources/test")
    workbook = load_workbook(f"{data_dir}/工作簿5.xlsx")
    sheet = workbook["Sheet1"]

    final_result = []

    service_result = QueryMysqlDatabase().query_test1()

    row_list = list(sheet.iter_rows())
    for row in tqdm(row_list, desc="获取信息中......"):
        # 获取客户id
        company_name = row[0].value
        service_name = row[1].value

        for service in service_result:
            if service['service_name'] == service_name and service['company_name'] == company_name:
                final_result.append(service)
                break


    # 创建新表
    new_workbook = Workbook()
    # 汇总表格
    new_sheet_one = new_workbook['Sheet']
    new_sheet_one.title = '汇总'
    new_sheet_one.append(['客户ID', '客户名称', '事项ID', '事项名称', '开始时间', '结束时间', '暂停时间', '交付时间', '法律顾问'])
    
    for row in final_result:
        row_data = [
            row["customer_id"],
            row["company_name"],
            row["service_id"],
            row["service_name"],
            row["start_time"],
            row["end_time"],
            row["service_stop"],
            row["delivery_time"],
            row["lawers"]
        ]
        new_sheet_one.append(row_data)

    new_workbook.save(
        f'{paths.get_resource_path('resources/test')}/事项数据.xlsx')


if __name__ == '__main__':
    test1()