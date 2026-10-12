from src.utils import paths_utils as paths
from openpyxl.reader.excel import load_workbook
from src.mapper.local.local_db_insert_impl import LocalMysqlDatabase as LocalInsertMysqlDatabase
from src.utils.dict_result import DictResult

"""批量导入退款客户，使用本地数据库存储"""
def public_batch_insert_refound(data):
    # 获取路径
    data_dir = paths.get_resource_path("resources/insert_refound_customer_excel_templates")

    # 获取文件
    workbook = load_workbook(f'{data_dir}/{data['refoundFileName']}.xlsx')
    sheet = workbook['Sheet1']

    # 获取数据
    sheet_data_list = list(sheet.iter_rows(min_row= 2, max_row= sheet.max_row))
    refound_list = [
        {
            # 到账时间（与系统一致）
            'accounting_time': row[1].value,
            # 退款主体
            'subject': row[2].value,
            # 审批单号
            'audit_number': row[3].value,
            # 分公司名称
            'salesman_company': row[4].value,
            # 业绩归属公司名称
            'performance_company': row[5].value,
            # 业务员名称
            'salesman_name': row[6].value,
            # 客户名称
            'company_name': row[7].value,
            # 收款金额
            'receive_money': row[8].value,
            # 是否已建群服务
            'is_group': 1 if row[9].value and row[9].value == '是' else 0,
            # 法律顾问
            'lawers': row[10].value,
            # 服务顾问
            'server': row[11].value,
            # 是否是续费
            'is_renewal': 1 if row[12].value and row[12].value == '是' else 0,
            # 合同开始时间
            'start_time': row[13].value,
            # 合同结束时间
            'end_time': row[14].value,
            # 套餐类型
            'criterion_type': row[15].value,
            # 退款时间
            'refound_time': row[16].value,
            # 退款金额
            'refound_money': row[17].value,
            # 是否终止
            'is_end': 1 if row[18].value and row[18].value == '是否终止' else 0,
            # 退款原因
            'refound_cause': row[19].value
        } for row in sheet_data_list if row[7].value is not None
    ]

    # 批量新增
    LocalInsertMysqlDatabase().batch_insert_refound(refound_list)

    return DictResult.success_response(
        message=f"新增{len(refound_list)}条退费客户成功",
        status_code=200,
        types="batch_insert_refound",
        data=None
    )